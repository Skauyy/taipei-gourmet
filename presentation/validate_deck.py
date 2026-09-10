#!/usr/bin/env python3
"""Validate OOXML and optionally render a contact sheet with local LibreOffice."""
import argparse
import io
import hashlib
import json
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET
from PIL import Image, ImageDraw
from pptx import Presentation

parser=argparse.ArgumentParser()
parser.add_argument('deck',type=Path)
parser.add_argument('--report',type=Path,required=True)
parser.add_argument('--preview',type=Path)
parser.add_argument('--render-dir',type=Path)
args=parser.parse_args()
ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
with zipfile.ZipFile(args.deck) as z:
    assert z.testzip() is None, 'Corrupt ZIP member'
    names=set(z.namelist())
    slides=sorted((n for n in names if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),key=lambda x:int(re.search(r'(\d+)\.xml',x)[1]))
    assert len(slides)==10, f'Expected ten slides, got {len(slides)}'
    for n in names:
        if n.endswith(('.xml','.rels')): ET.fromstring(z.read(n))
    missing=[]
    for n in names:
        if not n.endswith('.rels'): continue
        base=posixpath.dirname(posixpath.dirname(n))
        for rel in ET.fromstring(z.read(n)):
            if rel.attrib.get('TargetMode')=='External': continue
            target=rel.attrib['Target']
            resolved=target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(base,target))
            if resolved not in names: missing.append((n,target))
    assert not missing, missing
    images=[n for n in names if n.startswith('ppt/media/')]
    photo_manifest = json.loads((Path(__file__).resolve().parents[1] / 'data/restaurant-photos.json').read_text())
    allowed_hashes = {hashlib.sha256((Path(__file__).resolve().parents[1] / photo['src']).read_bytes()).hexdigest() for photo in photo_manifest.values()}
    for n in images:
        with Image.open(io.BytesIO(z.read(n))) as im: im.verify()
        assert hashlib.sha256(z.read(n)).hexdigest() in allowed_hashes, f'Unverified restaurant photograph: {n}'
    counts=[]
    for n in slides:
        root=ET.fromstring(z.read(n))
        counts.append({'slide':int(re.search(r'(\d+)\.xml',n)[1]),'text_runs':len(root.findall('.//a:t',ns)), 'pictures':len(root.findall('.//p:pic',ns))})
        assert counts[-1]['text_runs']>=8
    notes=[n for n in names if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n)]
    assert len(notes)==10
    for n in notes:
        note_text = ''.join(ET.fromstring(z.read(n)).itertext())
        assert '資料來源' in note_text and '必比登' in note_text and 'CC BY-SA 4.0' in note_text

prs=Presentation(args.deck)
assert abs(prs.slide_width/prs.slide_height-16/9)<.00001
outside=[]
for i,s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.left < -100 or sh.top < -100 or sh.left+sh.width>prs.slide_width+100 or sh.top+sh.height>prs.slide_height+100:
            outside.append((i+1,sh.name))
assert not outside, outside
report={'deck':args.deck.name,'slide_count':10,'aspect_ratio':'16:9','zip_crc':'pass','xml_parse':'pass','internal_relationships':'pass','missing_images':0,'embedded_image_files':len(images),'native_text_runs':sum(s['text_runs'] for s in counts),'picture_instances':sum(s['pictures'] for s in counts),'slides_with_notes':len(notes),'out_of_bounds_shapes':outside,'slides':counts}
data=json.loads((Path(__file__).parent/'guide-data.json').read_text())
report['guide_counts']={'entries':len(data['shops']),'kinds':dict(Counter(x['k'] for x in data['shops'])),'zones':dict(Counter(x['z'] for x in data['shops'])),'min_price':min(x['p'] for x in data['shops']),'max_price':max(x['p'] for x in data['shops'])}
assert report['guide_counts']['entries']==88
assert report['guide_counts']['zones']=={'arrival':13,'yt':41,'ntu':22,'other':12}
assert report['guide_counts']['kinds']=={'food':63,'drink':15,'dessert':10}
bib_data=json.loads((Path(__file__).resolve().parents[1]/'data/bib-gourmand.json').read_text())
report['bib_collection']={'records':len(bib_data['restaurants']),'current':sum(x['bib']['current'] for x in bib_data['restaurants']),'latest_edition':bib_data['latestEdition']}
report['venue_photo_allowlist']='pass'
if args.preview:
    import pypdfium2 as pdfium
    assert shutil.which('libreoffice'), 'Install LibreOffice Impress to render previews'
    render_dir=args.render_dir or Path(tempfile.mkdtemp(prefix='taipei-deck-'))
    render_dir.mkdir(parents=True,exist_ok=True)
    profile=(render_dir/'lo-profile').resolve().as_uri()
    subprocess.run(['libreoffice',f'-env:UserInstallation={profile}','--headless','--convert-to','pdf','--outdir',str(render_dir),str(args.deck.resolve())],check=True,timeout=120)
    pdf=pdfium.PdfDocument(render_dir/(args.deck.stem+'.pdf'))
    assert len(pdf)==10
    thumb_w,thumb_h=800,450
    sheet=Image.new('RGB',(1660,2460),'#D8D4C7')
    draw=ImageDraw.Draw(sheet)
    for i in range(len(pdf)):
        page=pdf[i]
        assert abs(page.get_width()/page.get_height()-16/9)<.001
        bitmap=page.render(scale=2)
        image=bitmap.to_pil()
        image.save(render_dir/f'slide-{i+1:02d}.png')
        x=20+(i%2)*820; y=20+(i//2)*488
        sheet.paste(image.resize((thumb_w,thumb_h),Image.Resampling.LANCZOS),(x,y))
        draw.text((x,y+457),f'{i+1:02d}  /  TAIWAN GOURMET',fill='#183E35')
        textpage=page.get_textpage(); extracted=textpage.get_text_range()
        assert len(extracted)>80
        assert '\ufffd' not in extracted
        (render_dir/f'slide-{i+1:02d}.txt').write_text(extracted)
        textpage.close();image.close();bitmap.close();page.close()
    pdf.close(); args.preview.parent.mkdir(parents=True,exist_ok=True);sheet.save(args.preview,optimize=True)
    report['render']={'engine':'LibreOffice Impress','rendered_pages':10,'preview':args.preview.name,'text_extraction':'pass','visual_review':'Contact sheet and full-resolution slides should be reviewed after each rebuild.'}
args.report.parent.mkdir(parents=True,exist_ok=True)
args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(report,ensure_ascii=False,indent=2))
