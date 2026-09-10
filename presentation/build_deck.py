#!/usr/bin/env python3
"""Build an editable editorial deck from the reviewed guide-data.json snapshot."""
import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
DATA = json.loads((HERE / 'guide-data.json').read_text())
SHOPS = {s['n']: s for s in DATA['shops']}
ASSETS = HERE.parent / 'assets'
BIB = json.loads((HERE.parent / 'data/bib-gourmand.json').read_text())
PHOTOS = json.loads((HERE.parent / 'data/restaurant-photos.json').read_text())
ALL_SHOPS = list(DATA['shops'])
for entry in BIB['restaurants']:
    if entry['n'] not in SHOPS and not any(name in SHOPS for name in entry.get('aliases', [])):
        ALL_SHOPS.append({'n': entry['n'], 'p': None, 'z': 'other', 'k': entry['k']})
CURRENT = sum(entry['bib']['current'] for entry in BIB['restaurants'])
PAST = len(BIB['restaurants']) - CURRENT
PHOTO_NOTES = '\n'.join(f"{name} ({photo['year']}): {photo['author']}; {photo['license']}; {photo['licenseUrl']}; {photo['sourceUrl']}; {photo['changes']}" for name, photo in PHOTOS.items())
IVORY, FOREST, RED, OLIVE, INK, RULE = 'F5F1E8', '183E35', 'DE4932', '777D59', '243A32', 'D8D4C7'
CN, SERIF, SANS = 'Noto Sans CJK TC', 'Liberation Serif', 'Inter'
W, H = 13.333333, 7.5
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Taiwan Gourmet / 台灣美食'
prs.core_properties.subject = 'An editorial food field guide for the Yuantong and NTU neighborhoods'
prs.core_properties.author = 'Taiwan Gourmet'
prs.core_properties.keywords = 'Taipei, 台灣美食, Traditional Chinese, editable, guide estimates'
COMMON = ('資料來源：原指南 app.js 店舖資料及 data/bib-gourmand.json 的台北、新北歷屆必比登核實紀錄。'
          '人均為原指南估算，不是單品現價；新增必比登未知價格及營業時間留空。'
          '必比登只屬來源所列店家，不自動延伸到同品牌分店；未列年份不代表未獲獎。'
          '出發前請向店家確認地址、營業安排及價格。照片只採可核實及授權的指定店家照片，均標記拍攝年份。\n' + PHOTO_NOTES)

def rgb(v): return RGBColor.from_string(v)
def rect(s, x,y,w,h,fill, radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.fill.background()
    if radius: sh.adjustments[0]=0.08
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in sh._element.xpath('./p:style/a:effectRef'): effect.set('idx', '0')
    return sh

def text(s, value,x,y,w,h,size=18,color=INK,font=CN,bold=False,tracking=None):
    box=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_before=Pt(0); p.space_after=Pt(0)
        p.line_spacing=1.12
        run=p.add_run(); run.text=line
        run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=rgb(color)
        rp=run._r.get_or_add_rPr()
        ea=OxmlElement('a:ea'); ea.set('typeface',CN); rp.append(ea)
        if tracking is not None: rp.set('spc',str(tracking))
    return box

def line(s,x,y,w,color=RULE):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x),Inches(y),Inches(x+w),Inches(y))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(.7)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in sh._element.xpath('./p:style/a:effectRef'): effect.set('idx', '0')

def photo(s,name,x,y,w,h):
    path=ASSETS/name
    assert path.is_file(), path
    sh=s.shapes.add_picture(str(path), Inches(x),Inches(y))
    sh.crop_left=sh.crop_right=sh.crop_top=sh.crop_bottom=0
    from PIL import Image
    iw,ih=Image.open(path).size
    ratio=(w/h)/(iw/ih)
    if ratio<1: sh.crop_left=sh.crop_right=(1-ratio)/2
    else: sh.crop_top=sh.crop_bottom=(1-1/ratio)/2
    sh.width,sh.height=Inches(w),Inches(h)
    return sh

def label(s,v,x,y,w=10,color=OLIVE): return text(s,v,x,y,w,.24,10,color,SANS,True,160)
def footer(s,num,dark=False,custom=None):
    c='B8C5B7' if dark else OLIVE
    line(s,.6,7.02,12.13,'527064' if dark else RULE)
    text(s,custom or '台灣美食・雙北版 / 人均為原指南估算；歷史入選及照片不代表即時店況',.6,7.12,11.5,.2,8.5,c)
    text(s,f'{num:02d}',12.2,7.1,.5,.25,10,c,SANS)

def slide(num,tag,title,subtitle=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(FOREST if dark else IVORY)
    label(s,tag,.6,.36,color='B8C5B7' if dark else OLIVE)
    text(s,title,.6,.87,12.1,.65,36,IVORY if dark else FOREST,SERIF)
    if subtitle: text(s,subtitle,.63,1.63,12,.55,17,IVORY if dark else INK)
    footer(s,num,dark)
    s.notes_slide.notes_text_frame.text=COMMON
    return s

def note(s,v): s.notes_slide.notes_text_frame.text=COMMON+'\n\n'+v

def shoprow(s,name,x,y,w,caption,location=None):
    d=SHOPS[name]
    text(s,name,x,y,w-1.5,.4,20,FOREST,bold=True)
    text(s,f"NT${d['p']}",x+w-1.42,y+.02,1.42,.35,18,RED,SANS,True)
    text(s,caption,x,y+.43,w,.31,14,INK)
    if location: text(s,location,x,y+.79,w,.25,11,OLIVE)
    line(s,x,y+1.08,w)

# 01 Cover
s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,W,H,FOREST)
photo(s,'restaurants/yongkang.jpg',6.2,0,7.1333,7.5)
label(s,'TAIWAN  /  TAIPEI + NEW TAIPEI',.65,.55,5.2,'C6CDBD')
line(s,.65,1.05,4.9,'527064')
text(s,'Taiwan\nGourmet',.6,1.53,5.65,2.2,65,IVORY,SERIF)
text(s,'台灣美食',.65,4.05,5.1,.7,34,IVORY,bold=True)
text(s,'由落機嗰餐，到落堂嗰碗。',.67,5.08,5,.5,19,'DCE0D0')
label(s,f"{len(ALL_SHOPS)} ENTRIES  ·  BIB GOURMAND NOTEBOOK",.67,6.26,5,'E5B5A1')
text(s,'台北・新北日常與歷屆必比登指南',.67,6.69,5.25,.35,11,'C6CDBD')
rect(s,7.9,6.87,4.95,.29,IVORY)
text(s,'永康牛肉麵 · 2006 / Minghong · CC BY-SA 4.0',8.0,6.92,4.78,.2,8.5,FOREST)
note(s,'封面為已核實永康牛肉麵 2006 年照片，非目前份量或菜單保證。作者 Minghong，採 CC BY-SA 4.0，已縮圖及裁切；來源與完整授權見本頁備註。')

# 02 Overview
s=slide(2,'01  /  THE GUIDE','Two cities, by appetite.','唔使一次食晒，先揀你今日想食嗰區。')
photo(s,'restaurants/lan-jia.jpg',8.5,2.35,4.2,4.16)
text(s,'藍家割包 · 2007 / 竹筍弟弟 · 自由使用授權',8.5,6.61,4.2,.22,9,OLIVE)
text(s,str(len(ALL_SHOPS)),.6,2.35,3.3,1.48,100,RED,SERIF)
text(s,'筆美食條目',.66,4.0,3.2,.45,24,FOREST,bold=True)
text(s,'店舖、攤位與美食街一齊收錄。',.67,4.66,7.1,.4,15,OLIVE)
text(s,'NT$40–600',4.08,2.66,4.05,.7,33,FOREST,SERIF)
text(s,'原指南人均估算範圍',4.1,3.51,4,.4,16,INK)
text(s,'新增條目未核實的價格留空',4.1,4.03,4,.4,14,OLIVE)
line(s,.65,5.31,7.2)
for i,(num,name) in enumerate([(CURRENT,'現屆'),(PAST,'非現屆'),(len(PHOTOS),'真相片')]):
    x=.65+i*2.47
    text(s,str(num),x,5.55,1.45,.75,43,FOREST,SERIF)
    text(s,name,x+1.15,5.89,1.15,.4,16,INK)
note(s,f"總條目 {len(ALL_SHOPS)}，原指南88筆加上未重複的必比登條目。2026現屆 {CURRENT} 家；曾入選但非現屆 {PAST} 家。相片 {len(PHOTOS)} 張逐店核實，其他店不使用示意照片代替。")

# 03 Zones
s=slide(3,'02  /  FOUR ZONES','Find your food neighborhood.','四個指南分區；價格範圍只計原指南已有估算的條目。')
zone_data=[('arrival','Arrival','到埗美食','桃園機場\n台北車站沿線','落機唔使急，\n先食一餐再安頓。'),('yt','Yuantong','圓通宿舍美食','工專路 · 華新街\n南勢角 · 興南夜市','早餐、便當、宵夜，\n慢慢建立日常飯堂。'),('ntu','Gongguan / NTU','台大美食','公館 · 溫州街\n汀州路 · 台大校區','落堂搵飯食，\n順便加杯嘢飲。'),('other','Taipei classics','雙北其他美食','雙北必比登與其他條目\n市區小吃與夜市','留返一日，\n食吓台北經典。')]
for i,(z,en,zh,loc,blurb) in enumerate(zone_data):
    x=.6+3.08*i; rows=[d for d in ALL_SHOPS if d['z']==z]; prices=[d['p'] for d in rows if d['p'] is not None]
    rect(s,x,2.47,2.88,4.12,'EBE8DC' if i%2==0 else 'E3E7DB')
    text(s,f'0{i+1}',x+.2,2.65,2,.9,53,RED,SERIF)
    label(s,en.upper(),x+.2,3.65,2.5,OLIVE)
    text(s,zh,x+.2,4.08,2.5,.4,21,FOREST,bold=True)
    text(s,f"{len(rows)} 筆  /  NT${min(prices)}–{max(prices)}",x+.2,4.63,2.5,.3,12,RED)
    text(s,loc,x+.2,5.17,2.5,.7,13,INK)
    text(s,blurb,x+.2,6.0,2.5,.48,11,OLIVE)
note(s,'生活圈沿用原指南四區，新增雙北必比登列入其他區；必比登與城市可獨立篩選。此頁價格範圍只涵蓋原指南已有人均數值者，不含新增未核價店家。')

# 04 Arrival
s=slide(4,'03  /  ARRIVAL','Land. Eat. Settle in.','到埗先搞掂肚餓，唔使第一日就趕行程。')
rect(s,8.25,2.35,4.45,3.68,FOREST)
text(s,'先安頓，\n再覓食。',8.65,2.9,3.65,1.8,42,IVORY,bold=True)
text(s,'第一餐，唔使趕。',8.65,5.1,3.65,.4,18,'DCE0D0')
text(s,'機場分店不自動沿用品牌其他店的必比登獎項。',8.25,6.15,4.45,.23,9,OLIVE)
text(s,'交通路線、票價與所需時間\n請出發前查官方最新資訊。',8.25,6.48,4.45,.45,11,OLIVE)
steps=[('01','落機後，按航廈揀一餐','指南例子：小王煮瓜 桃園機場 T2','滷肉飯／焢肉飯 · 人均估算 NT$120'),('02','如經台北車站，再補給','指南例子：台鐵便當 台北車站','排骨便當／雞腿便當 · 人均估算 NT$100'),('03','先安頓行李，再探索附近','圓通宿舍 → 工專路／華新街','按精神同肚餓程度揀，唔一定每站都食。')]
for i,(n,t,a,b) in enumerate(steps):
    y=2.53+i*1.43
    text(s,n,.65,y,.75,.58,29,RED,SERIF)
    text(s,t,1.55,y,6.3,.45,22,FOREST,bold=True)
    text(s,a,1.55,y+.55,6.3,.36,14,INK)
    text(s,b,1.55,y+.95,6.3,.3,12,OLIVE)
    if i<2: line(s,1.55,y+1.29,6.05)
note(s,'此頁為編輯建議順序，不保證途經台北車站，亦不描述可通行的航廈區域或機場店家現況。價格取自小王煮瓜 桃園機場 T2（p=120）、台鐵便當 台北車站（p=100）。無加入未经核實之車費或行車時間。')

# 05 Dorm
s=slide(5,'04  /  YUANTONG','Your everyday table.','宿舍附近，由早餐食到宵夜。')
rect(s,.6,2.36,4.65,3.83,FOREST)
text(s,'巷口日常',1.0,2.85,3.85,.75,35,IVORY,bold=True)
text(s,'工專路\n華新街\n南勢角',1.0,4.0,3.8,1.6,25,'DCE0D0')
text(s,'未核實店家照片的條目，以純文字介紹。',.6,6.3,4.65,.22,9,OLIVE)
label(s,'41 ENTRIES  /  NT$40–320',.6,6.65,4.7,RED)
for i,args in enumerate([
 ('早貓（早點）','蛋餅做起點，留返午餐再慢慢揀。','工專路 20 號'),
 ('小舖媽','小吃、義大利麵、自助餐，日常一餐。','工專路 26 號'),
 ('小檳城食堂','燒雞飯、炒粿條、叻沙，換吓南洋口味。','華新街 112 號'),
 ('口碑鹹酥雞','鹹酥雞、魷魚、雞皮，宵夜可以分享。','忠孝街 76 號')]):
    shoprow(s,args[0],5.78,2.33+1.14*i,6.92,args[1],args[2])
note(s,'四筆精選均來自 z=yt，p 依次 60、90、170、85。全部有 dc=1 標記；該標記來自指南對 Dcard 文章及留言的整理，未重新核實作者體驗或店家現況。地址為原資料摘錄。')

# 06 NTU
s=slide(6,'05  /  GONGGUAN + NTU','Between classes, eat well.','公館搵嘢食，正餐、小食、甜品都留一格。')
photo(s,'restaurants/lan-jia.jpg',.6,2.35,4.65,3.85)
text(s,'照片只屬藍家割包 · 2007 / 竹筍弟弟 · 自由使用授權',.6,6.35,4.8,.24,9,OLIVE)
label(s,'22 ENTRIES  /  NT$50–400',.6,6.66,4.7,RED)
for i,args in enumerate([
 ('JODO 飯糰','飯糰、鮪魚飯糰，簡單開始一日。','汀州路三段 131 號'),
 ('藍家割包','瘦肉割包、八寶麵，先揀一樣試吓。','羅斯福路三段 316 巷 8 弄 3 號'),
 ('池先生 Kopitiam','海南雞飯、叻沙，想食正餐就揀呢邊。','羅斯福路三段 284 巷 10 號'),
 ('鴉片粉圓','粉圓冰、綜合冰，食完飯加個甜品。','羅斯福路四段 52 巷 16 弄 4 號')]):
    shoprow(s,args[0],5.78,2.33+1.14*i,6.92,args[1],args[2])
note(s,'四筆精選均來自 z=ntu，p 依次 60、90、170、65。餐點介紹以 s 招牌欄為依據，語氣為編輯改寫；並非作者實地試食記錄。照片來源明確指向公館藍家割包，攝於2007年；不代表其他三間店，也不是現在的菜單或份量保證。')

# 07 Bib Gourmand notebook
s=slide(7,'06  /  THE BIB GOURMAND NOTEBOOK','Good food, with a paper trail.',f"雙北 {CURRENT} 家現屆入選，{PAST} 家曾入選但非現屆；每個年份都有來源。")
for i,name in enumerate(['阜杭豆漿','永康牛肉麵','藍家割包']):
    x=.6+4.1*i
    entry=next(row for row in BIB['restaurants'] if name in row.get('aliases',[]) or row['n']==name)
    photo_data=PHOTOS[name]
    photo(s,photo_data['src'].removeprefix('assets/'),x,2.35,3.9,2.35)
    text(s,f"{photo_data['year']} 年照片 / {photo_data['author']}",x,4.79,3.9,.2,8.5,OLIVE)
    label(s,'2026 BIB GOURMAND' if entry['bib']['current'] else 'PAST BIB GOURMAND',x,5.19,3.9,RED)
    text(s,name,x,5.6,3.9,.45,23,FOREST,bold=True)
    text(s,f"已核實 {len(entry['bib']['years'])} 個年份\n逐年來源與照片授權見本頁備註。",x,6.15,3.9,.56,12,INK)
    s.notes_slide.notes_text_frame.text += '\n'+name+': '+json.dumps(entry['bib'],ensure_ascii=False)
s.notes_slide.notes_text_frame.text += '\n當年必比登並非今日營業保證；沒有把任意品牌分店當成獲獎店家。'

# 08 Editorial itinerary
s=slide(8,'07  /  AN EDITORIAL DAY','One day. Four good pauses.','編輯提案：圓通 → 公館 → 圓通，唔係既定或已實測路線。',True)
line(s,.9,3.0,11.55,'688274')
for i,(time,name,food) in enumerate([('早餐','早貓（早點）','蛋餅'),('午餐','池先生 Kopitiam','海南雞飯／叻沙'),('午後','鴉片粉圓','粉圓冰／綜合冰'),('夜晚','口碑鹹酥雞','鹹酥雞，按食量加減')]):
    x=.6+3.08*i
    rect(s,x,2.77,.5,.5,RED)
    text(s,str(i+1),x+.16,2.84,.22,.23,13,IVORY,SANS,True)
    label(s,time,x,3.6,2.8,'C6CDBD')
    text(s,name,x,4.1,2.9,.8,23,IVORY,bold=True)
    text(s,food,x,5.02,2.86,.45,13,'C6CDBD')
    text(s,f"NT${SHOPS[name]['p']}",x,5.6,2.9,.56,29,'F0B39E',SERIF)
rect(s,.6,6.4,12.1,.42,'294B40')
text(s,'四站人均估算合計約 NT$380；不含交通、加點。先核對休息日，再決定行唔行。',.8,6.48,11.7,.26,12,IVORY)
note(s,'本頁明確為編輯建議，不是已核實旅程、交通導航或營業承諾。60+170+65+85=380，僅加總四筆 p 人均估算，並非指定單品報價。原資料中早貓、口碑鹹酥雞週日休；因此不應假定所有日期均可照行。')

# 09 Site use
s=slide(9,'08  /  USE THE WEBSITE','Less scrolling. More eating.','先縮細清單，再打開地圖確認；唔好淨係睇個「營業中」。')
rect(s,.6,2.4,5.28,4.37,'E7E9DE')
label(s,'A SIMPLE WAY TO CHOOSE',.85,2.68,4.8,OLIVE)
for j,(v,active) in enumerate([('歷屆必比登',True),('新北',False),('入選年份與來源 ↗',False)]):
    rect(s,.88,3.18+j*.76,4.68,.55,FOREST if active else IVORY,True)
    text(s,v,1.08,3.29+j*.76,4.2,.3,17,IVORY if active else FOREST)
text(s,'介面流程示意 · 先縮細清單再揀',.88,5.66,4.7,.33,11,OLIVE)
text(s,'資料未定 ≠ 一定休息\n有結果 ≠ 保證今日有開',.88,6.06,4.7,.58,15,FOREST)
items=[('01  歷屆 × 現屆','必比登、城市、生活圈與種類可交叉篩選。'),('02  年份 × 來源','每個入選年份有出處；不把品牌獎項套給分店。'),('03  依家營業緊？','按資料內時間推算，非店家即時回報；出門再確認。'),('04  生活圈地圖','點選聚落睇推介，再確認實際位置與交通。')]
for i,(a,b) in enumerate(items):
    y=2.49+i*1.05
    text(s,a,6.38,y,6.25,.4,20,FOREST,bold=True)
    text(s,b,6.38,y+.5,6.25,.5,13,INK)
    if i<3:line(s,6.38,y+.9,6.25)
note(s,'功能依 app.js 核對：zone/kind 多選、p/r 雙向排序、openState 以台北時間依靜態 o 欄計算、地址 mapsUrl Google Maps 搜尋連結、生活圈 SPOTS 示意地圖。搜尋對應店名 n、地址 a、招牌 s。未見店舖收藏，不宣稱存在此功能。不展示星分，避免誤認為 Google 評分。')

# 10 Closing
s=slide(10,'09  /  BEFORE YOU GO','Eat curious. Check first.','食之前做少少功課，食嘅時候就慢慢享受。')
rect(s,.6,2.43,4.43,4.3,FOREST)
text(s,'一份指南，\n唔係保證。',.93,2.85,3.85,1.57,36,IVORY,bold=True)
text(s,'留返肚餓，\n畀下一條街。',.95,5.22,3.65,.99,25,'D7DECE')
for i,(a,b) in enumerate([
 ('出發前再確認','地址、休息日、供應狀況同價錢，以店家最新公告為準。'),
 ('推薦指數要識睇','屬指南主觀參考，唔係即時 Google 評分或品質保證。'),
 ('照片要對得上店家','只用有店家身分證據及授權的照片；其餘保留文字。')]):
    y=2.43+i*1.06
    text(s,a,5.64,y,6.95,.4,20,FOREST,bold=True)
    text(s,b,5.64,y+.47,6.95,.49,13,INK)
line(s,5.64,5.69,6.95)
label(s,'SOURCE CONTEXT',5.64,5.95,6.9,RED)
text(s,'原指南88筆 + 雙北必比登核實資料。\n歷屆年份來自米其林年度名單及具日期的完整報導。\n相片署名、拍攝年份、授權與原始網址均列於各頁備註。',5.64,6.29,7.0,.65,10.5,OLIVE)
note(s,'原指南部分宿舍建議來自Dcard宿舍攻略及留言，未重新試食。必比登原始年度來源見 data/bib-gourmand.json；相片身分、作者、授權與修改方式見 data/restaurant-photos.json。')

assert len(prs.slides)==10
parser=argparse.ArgumentParser()
parser.add_argument('--output',type=Path,default=HERE/'Taiwan-Gourmet.pptx')
args=parser.parse_args(); args.output.parent.mkdir(parents=True,exist_ok=True)
prs.save(args.output)
print(f'Created {args.output}: {len(prs.slides)} editable 16:9 slides')
